"""Excel / CSV 导入 — SRS-IMP-001~003 / AC-IMP-01。

预览与提交严格分离:

  预览阶段解析全部行、逐行校验、把结果连同原始数据一起落库(import_batch_row),
  但不产生任何业务数据。使用者看到的是"第 N 行为什么不行", 不是一句"导入失败"。

  提交阶段只有在 error_rows = 0 时才允许 —— 数据库上还有
  ck_import_batch_commit 兜底。这里刻意不做"跳过错误行继续导入": 部分成功会留下
  一个谁也说不清完整性的 BOM, 排查成本远高于让使用者改完再传一次。

导入不绕过任何业务规则。BOM 行仍走 bom.add_line, 因此自引用、循环、数量非正
一样会被数据库触发器拦下 —— 批量入口是数据规则最容易被绕过的地方, 必须走同一条路。
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from typing import Any

import psycopg

from .. import audit
from ..db import execute, fetch_all, fetch_one, scalar
from . import bom as bom_svc

# BOM 导入模板列名 → 内部字段。中英文表头都接受, 因为现场表格两种都有。
BOM_COLUMNS = {
    "item_number": ("项号", "item", "item_number", "序号"),
    "child_object_code": ("子件号", "子项", "child", "part_number", "p/n", "件号"),
    "quantity": ("数量", "qty", "quantity", "用量"),
    "unit_code": ("单位", "unit", "uom"),
    "reference_designator": ("位号", "designator", "reference_designator"),
    "effectivity": ("有效性", "effectivity"),
    "applicability_rule_code": ("适用性规则", "applicability", "applicability_rule", "rule_code"),
    "notes": ("备注", "notes", "remark"),
}
REQUIRED = ("item_number", "child_object_code", "quantity")


def _norm(s: Any) -> str:
    return str(s or "").strip().lower().replace(" ", "").replace("　", "")


def _map_header(header: list[str]) -> tuple[dict[int, str], list[str]]:
    """把表头列映射到内部字段, 返回 (列序号→字段, 缺失的必填列)。"""
    mapping: dict[int, str] = {}
    for idx, col in enumerate(header):
        n = _norm(col)
        for field, aliases in BOM_COLUMNS.items():
            if n in {_norm(a) for a in aliases}:
                mapping[idx] = field
                break
    missing = [f for f in REQUIRED if f not in mapping.values()]
    return mapping, missing


def parse_table(content: bytes, filename: str) -> tuple[list[str], list[list[Any]]]:
    """解析 xlsx 或 csv, 返回 (表头, 数据行)。"""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        text = content.decode("utf-8-sig", errors="replace")
        rows = [r for r in csv.reader(io.StringIO(text))]
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return [], []
    return [str(c or "") for c in rows[0]], rows[1:]


# ---------------------------------------------------------------------
# 预览
# ---------------------------------------------------------------------
def preview_bom(conn: psycopg.Connection, parent_object_code: str, content: bytes,
                filename: str, actor: dict) -> dict:
    parent = fetch_one(conn, """
        SELECT id, object_code, display_name, lifecycle_status
          FROM design_object WHERE object_code = %s
    """, (parent_object_code,))
    if parent is None:
        raise LookupError(f"父项对象不存在: {parent_object_code}")

    header, data_rows = parse_table(content, filename)
    mapping, missing = _map_header(header)

    seq = scalar(conn, """
        SELECT COALESCE(max(substring(batch_number from 5)::int), 0) + 1
          FROM import_batch WHERE batch_number LIKE 'IMP-%'
    """) or 1
    batch = fetch_one(conn, """
        INSERT INTO import_batch (batch_number, import_type, source_filename,
                                  source_sha256, target_context, created_by)
        VALUES (%s, 'BOM', %s, %s, %s::jsonb, %s)
        RETURNING id, batch_number, status
    """, (f"IMP-{seq:06d}", filename, hashlib.sha256(content).hexdigest(),
          json.dumps({"parent_object_code": parent_object_code}, ensure_ascii=False),
          actor["user_id"]))

    if missing:
        # 表头就不对时不再逐行报错 —— 那只会刷出几百条同样的消息
        execute(conn, """
            UPDATE import_batch SET total_rows = %s, error_rows = 1 WHERE id = %s
        """, (len(data_rows), batch["id"]))
        execute(conn, """
            INSERT INTO import_batch_row (import_batch_id, row_number, raw_data, result, messages)
            VALUES (%s, 0, %s::jsonb, 'ERROR', %s::jsonb)
        """, (batch["id"], json.dumps({"header": header}, ensure_ascii=False),
              json.dumps([f"缺少必需列: {', '.join(missing)}; "
                          f"可接受的列名见导入模板"], ensure_ascii=False)))
        return _batch_result(conn, str(batch["id"]))

    seen_items: set[tuple[str,str]] = set()
    ok = warn = err = 0

    for i, raw in enumerate(data_rows, start=2):     # 第 1 行是表头, 数据从第 2 行起
        rec: dict[str, Any] = {}
        for idx, field in mapping.items():
            rec[field] = raw[idx] if idx < len(raw) else None

        messages: list[str] = []
        result = "OK"

        # 必填
        for f in REQUIRED:
            if rec.get(f) in (None, ""):
                messages.append(f"必填项为空: {f}")
                result = "ERROR"

        # 数量
        qty = None
        if rec.get("quantity") not in (None, ""):
            try:
                qty = float(str(rec["quantity"]).strip())
                if qty <= 0:
                    messages.append(f"数量必须大于 0, 实际 {qty} (CK-02)")
                    result = "ERROR"
            except ValueError:
                messages.append(f"数量不是有效数字: {rec['quantity']!r}")
                result = "ERROR"

        # 子项存在性
        code = str(rec.get("child_object_code") or "").strip()
        child = None
        if code:
            child = fetch_one(conn, """
                SELECT id, object_code, lifecycle_status FROM design_object
                 WHERE object_code = %s
            """, (code,))
            if child is None:
                # 试试交叉引用归一匹配 —— 现场表格里常填历史件号
                alt = fetch_one(conn, """
                    SELECT d.id, d.object_code, d.lifecycle_status
                      FROM cross_reference cr JOIN design_object d ON d.id = cr.design_object_id
                     WHERE cr.normalized_value = dcms_normalize_identifier(%s) LIMIT 1
                """, (code,))
                if alt is not None:
                    child = alt
                    messages.append(f"按交叉引用匹配到 {alt['object_code']} "
                                    f"(原值 {code}), 请确认")
                    result = "WARNING" if result == "OK" else result
                else:
                    messages.append(f"子项对象不存在: {code}")
                    result = "ERROR"
            if child is not None and str(child["id"]) == str(parent["id"]):
                messages.append("子项与父项相同, BOM 不得自引用 (INV-008)")
                result = "ERROR"

        # 同一 ITEM 可有不同替代件（由 Applicability 区分）；同 ITEM + 同子件仍视为重复。
        item = str(rec.get("item_number") or "").strip()
        child_code = str(rec.get("child_object_code") or "").strip()
        if item and child_code:
            key = (item, child_code)
            if key in seen_items:
                messages.append(f"项号 {item} / 子件 {child_code} 在本文件中重复")
                result = "ERROR"
            seen_items.add(key)

        rule_code = str(rec.get("applicability_rule_code") or "").strip()
        if rule_code and not scalar(conn, "SELECT EXISTS(SELECT 1 FROM applicability_rule WHERE rule_code=%s AND status='ACTIVE')", (rule_code,)):
            messages.append(f"Applicability 规则不存在或不可用: {rule_code}")
            result = "ERROR"

        # 单位
        unit = str(rec.get("unit_code") or "").strip() or None
        if unit and not scalar(conn, "SELECT EXISTS(SELECT 1 FROM unit WHERE code=%s)", (unit,)):
            messages.append(f"单位 {unit} 不在受控单位字典中")
            result = "ERROR"

        # 草稿子项只是提示
        if child is not None and child["lifecycle_status"] == "DRAFT" and result == "OK":
            messages.append(f"子项 {child['object_code']} 仍为草稿状态 (DQ-BOM-002)")
            result = "WARNING"

        if result == "OK":
            ok += 1
        elif result == "WARNING":
            warn += 1
        else:
            err += 1

        execute(conn, """
            INSERT INTO import_batch_row (import_batch_id, row_number, raw_data, result, messages)
            VALUES (%s,%s,%s::jsonb,%s,%s::jsonb)
        """, (batch["id"], i,
              json.dumps({k: (str(v) if v is not None else None) for k, v in rec.items()},
                         ensure_ascii=False),
              result, json.dumps(messages, ensure_ascii=False)))

    execute(conn, """
        UPDATE import_batch SET total_rows=%s, ok_rows=%s, warning_rows=%s, error_rows=%s
         WHERE id = %s
    """, (len(data_rows), ok, warn, err, batch["id"]))

    audit.write(conn, action="IMPORT_PREVIEW", user_id=str(actor["user_id"]),
                username=actor["username"], object_type="IMPORT_BATCH",
                object_id=str(batch["id"]), object_code=batch["batch_number"],
                new_value={"filename": filename, "parent": parent_object_code,
                           "total": len(data_rows), "ok": ok,
                           "warning": warn, "error": err},
                session_id=str(actor.get("session_id")), client_ip=actor.get("client_ip"))
    return _batch_result(conn, str(batch["id"]))


def _batch_result(conn: psycopg.Connection, batch_id: str) -> dict:
    b = fetch_one(conn, "SELECT * FROM import_batch WHERE id = %s", (batch_id,))
    rows = fetch_all(conn, """
        SELECT row_number, raw_data, result, messages, created_object_type, created_object_id
          FROM import_batch_row WHERE import_batch_id = %s ORDER BY row_number
    """, (batch_id,))
    return {
        "batch_id": str(b["id"]), "batch_number": b["batch_number"],
        "status": b["status"], "import_type": b["import_type"],
        "source_filename": b["source_filename"],
        "total_rows": b["total_rows"], "ok_rows": b["ok_rows"],
        "warning_rows": b["warning_rows"], "error_rows": b["error_rows"],
        "committable": b["status"] == "PREVIEW" and b["error_rows"] == 0,
        "rows": rows,
    }


# ---------------------------------------------------------------------
# 提交
# ---------------------------------------------------------------------
def commit_bom(conn: psycopg.Connection, batch_id: str, actor: dict,
               replace_existing: bool = False) -> dict:
    b = fetch_one(conn, "SELECT * FROM import_batch WHERE id = %s", (batch_id,))
    if b is None:
        raise LookupError("导入批次不存在")
    if b["status"] != "PREVIEW":
        raise ValueError(f"批次状态为 {b['status']}, 只有 PREVIEW 可提交")
    if b["error_rows"] > 0:
        raise ValueError(f"批次存在 {b['error_rows']} 个错误行, 不允许部分导入。"
                         f"请修正后重新上传")

    ctx = b["target_context"] or {}
    parent = fetch_one(conn, "SELECT id, object_code FROM design_object WHERE object_code=%s",
                       (ctx.get("parent_object_code"),))
    if parent is None:
        raise LookupError("父项对象不存在")

    header = bom_svc.get_or_create_working(conn, str(parent["id"]), actor)
    if replace_existing:
        execute(conn, "DELETE FROM bom_line WHERE bom_header_id = %s", (header["id"],))

    rows = fetch_all(conn, """
        SELECT id, row_number, raw_data FROM import_batch_row
         WHERE import_batch_id = %s AND result IN ('OK','WARNING') ORDER BY row_number
    """, (batch_id,))

    created = 0
    for r in rows:
        d = r["raw_data"]
        # 走与手工录入完全相同的入口, 数据库触发器一样会校验循环与自引用
        line = bom_svc.add_line(
            conn, str(parent["id"]),
            item_number=str(d["item_number"]).strip(),
            child_object_code=str(d["child_object_code"]).strip(),
            quantity=float(d["quantity"]),
            unit_code=(d.get("unit_code") or None),
            reference_designator=(d.get("reference_designator") or None),
            effectivity=(d.get("effectivity") or None),
            notes=(d.get("notes") or None), actor=actor)
        rule_code = (d.get("applicability_rule_code") or "").strip() if isinstance(d.get("applicability_rule_code"), str) else d.get("applicability_rule_code")
        if rule_code:
            from . import applicability as app_svc
            app_svc.assign_rule(conn, str(line["id"]), str(rule_code), actor)
        execute(conn, """
            UPDATE import_batch_row SET created_object_type='BOM_LINE', created_object_id=%s
             WHERE id = %s
        """, (line["id"], r["id"]))
        created += 1

    execute(conn, """
        UPDATE import_batch SET status='COMMITTED', committed_at=now(), committed_by=%s
         WHERE id = %s
    """, (actor["user_id"], batch_id))

    audit.write(conn, action="IMPORT_COMMIT", user_id=str(actor["user_id"]),
                username=actor["username"], object_type="IMPORT_BATCH",
                object_id=batch_id, object_code=b["batch_number"],
                new_value={"parent": parent["object_code"], "lines_created": created,
                           "replace_existing": replace_existing},
                session_id=str(actor.get("session_id")), client_ip=actor.get("client_ip"))
    return {"batch_number": b["batch_number"], "lines_created": created,
            "parent_object_code": parent["object_code"]}


def abort_batch(conn: psycopg.Connection, batch_id: str, reason: str, actor: dict) -> None:
    b = fetch_one(conn, "SELECT batch_number, status FROM import_batch WHERE id=%s", (batch_id,))
    if b is None:
        raise LookupError("导入批次不存在")
    if b["status"] != "PREVIEW":
        raise ValueError(f"批次状态为 {b['status']}, 不可放弃")
    execute(conn, "UPDATE import_batch SET status='ABORTED' WHERE id=%s", (batch_id,))
    audit.write(conn, action="IMPORT_ABORT", user_id=str(actor["user_id"]),
                username=actor["username"], object_type="IMPORT_BATCH",
                object_id=batch_id, object_code=b["batch_number"], reason=reason,
                session_id=str(actor.get("session_id")), client_ip=actor.get("client_ip"))


def bom_template_csv() -> str:
    """导入模板。给出可直接另存为 Excel 的表头与一行示例。"""
    return ("项号,子件号,数量,单位,位号,适用性规则,有效性,备注\r\n"
            "010,UG100001-001,2,EA,C1;C2,,全部,示例行——请删除后填写实际数据\r\n")
