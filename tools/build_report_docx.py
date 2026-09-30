"""把 UG-RPT 系列的 markdown 源生成为体系文件式样的 Word（页眉含 logo、编号、版本、日期）。

用法:
    python tools/build_report_docx.py <md> <docx> <标题> <编号> <版本> <日期>

生成后还需用 Word 更新目录域并导出 PDF。注意顺序:
先让 Word 保存一次(更新目录域的页码) → 再改标题颜色 → 最后只导 PDF 不回存。
因为 Word 保存时会按 themeColor 把内置标题样式的颜色刷回蓝色, 颜色必须改在最后。
"""
import os
import re
import shutil
import subprocess
import sys
import zipfile


from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

if len(sys.argv) != 7:
    sys.exit(__doc__)
MD, F, TITLE, NUM, REV, DATE = sys.argv[1:7]
LOGO = os.environ.get('DAS_LOGO', r'D:\UD-DCMS\设计保证体系文件\_模板\logo.png')

subprocess.run(['pandoc', MD, '-o', F, '--from', 'gfm', '--to', 'docx',
                '--toc', '--toc-depth=2'], check=True)

d = Document(F)
s0 = d.sections[0]
s0.page_width, s0.page_height = Cm(21), Cm(29.7)
s0.top_margin, s0.bottom_margin = Cm(3.0), Cm(2.1)
s0.left_margin = s0.right_margin = Cm(2.29)
s0.header_distance, s0.footer_distance = Cm(0.99), Cm(1.25)


def east(run, name, size=None, bold=None):
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    rf.set(qn('w:eastAsia'), name); rf.set(qn('w:ascii'), 'Arial'); rf.set(qn('w:hAnsi'), 'Arial')
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


HSIZE = {1: 15, 2: 13, 3: 11.5}
for p in d.paragraphs:
    st = (p.style.name or '') if p.style is not None else ''
    if st.startswith('Heading'):
        lvl = int(st[-1]) if st[-1].isdigit() else 2
        for r in p.runs:
            east(r, '黑体', HSIZE.get(lvl, 11.5), True)
    elif st.startswith('Title'):
        for r in p.runs:
            east(r, '黑体', 19, True)
    else:
        for r in p.runs:
            east(r, '宋体', 10.5)

for t in d.tables:
    try:
        t.style = d.styles['Table Grid']
    except Exception:
        pass
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    east(r, '宋体', 9, True if i == 0 else None)
            if i == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'DCE6F1')
                c._tc.get_or_add_tcPr().append(sh)
d.save(F)

# ---- 页眉页脚（与体系文件同一套式样）----
HDR = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
       '<w:tbl><w:tblPr><w:tblW w:w="9306" w:type="dxa"/><w:tblBorders>'
       '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="auto"/></w:tblBorders></w:tblPr>'
       '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="3800"/><w:gridCol w:w="3506"/></w:tblGrid>'
       '<w:tr><w:tc><w:tcPr><w:tcW w:w="2000" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr>'
       '<w:p><w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
       '<wp:extent cx="1200150" cy="342900"/><wp:docPr id="1" name="logo"/>'
       '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
       '<pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="logo.png"/><pic:cNvPicPr/></pic:nvPicPr>'
       '<pic:blipFill><a:blip r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
       '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1200150" cy="342900"/></a:xfrm>'
       '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
       '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p></w:tc>'
       '<w:tc><w:tcPr><w:tcW w:w="3800" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr>'
       '<w:p><w:pPr><w:jc w:val="center"/><w:spacing w:line="260" w:lineRule="auto"/></w:pPr>'
       '<w:r><w:rPr><w:rFonts w:eastAsia="黑体" w:ascii="Arial" w:hAnsi="Arial"/><w:b/>'
       '<w:sz w:val="21"/><w:szCs w:val="21"/></w:rPr><w:t>%s</w:t></w:r></w:p></w:tc>'
       '<w:tc><w:tcPr><w:tcW w:w="3506" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr>'
       '%s</w:tc></w:tr></w:tbl><w:p><w:pPr><w:spacing w:after="0"/></w:pPr></w:p></w:hdr>')


def line(t):
    return ('<w:p><w:pPr><w:spacing w:line="260" w:lineRule="auto" w:after="0"/></w:pPr>'
            '<w:r><w:rPr><w:rFonts w:eastAsia="黑体" w:ascii="Arial" w:hAnsi="Arial"/><w:b/>'
            '<w:sz w:val="21"/><w:szCs w:val="21"/></w:rPr><w:t>%s</w:t></w:r></w:p>' % t)


hdr = HDR % (TITLE, line('编号/No.: ' + NUM) + line('版本/Rev: ' + REV) + line('日期/Date: ' + DATE))
ftr = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
       '<w:p><w:pPr><w:jc w:val="right"/></w:pPr>'
       '<w:r><w:rPr><w:rFonts w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr><w:t>第</w:t></w:r>'
       '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
       '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
       '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
       '<w:r><w:rPr><w:rFonts w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr><w:t>页/共</w:t></w:r>'
       '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
       '<w:r><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>'
       '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
       '<w:r><w:rPr><w:rFonts w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr><w:t>页</w:t></w:r>'
       '</w:p></w:ftr>')
HDR_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/image" Target="media/logo.png"/></Relationships>')

z = zipfile.ZipFile(F); names = z.namelist(); parts = {k: z.read(k) for k in names}; z.close()
parts['word/header1.xml'] = hdr.encode('utf-8')
parts['word/footer1.xml'] = ftr.encode('utf-8')
parts['word/_rels/header1.xml.rels'] = HDR_RELS.encode('utf-8')
parts['word/media/logo.png'] = open(LOGO, 'rb').read()

ct = parts['[Content_Types].xml'].decode('utf-8')
if 'Extension="png"' not in ct:
    ct = ct.replace('<Default Extension="xml"',
                    '<Default Extension="png" ContentType="image/png"/><Default Extension="xml"', 1)
for nm, typ in (('header1', 'header'), ('footer1', 'footer')):
    tag = ('<Override PartName="/word/%s.xml" ContentType="application/vnd.openxmlformats-'
           'officedocument.wordprocessingml.%s+xml"/>' % (nm, typ))
    if tag not in ct:
        ct = ct.replace('</Types>', tag + '</Types>')
parts['[Content_Types].xml'] = ct.encode('utf-8')

rels = parts['word/_rels/document.xml.rels'].decode('utf-8')
used = {int(x) for x in re.findall(r'Id="rId(\d+)"', rels)}
hid, fid = max(used) + 1, max(used) + 2
rels = rels.replace('</Relationships>',
                    '<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/officeDocument/'
                    '2006/relationships/header" Target="header1.xml"/>'
                    '<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/officeDocument/'
                    '2006/relationships/footer" Target="footer1.xml"/></Relationships>' % (hid, fid))
parts['word/_rels/document.xml.rels'] = rels.encode('utf-8')

doc = parts['word/document.xml'].decode('utf-8')
doc = re.sub(r'(<w:sectPr[^>]*>)',
             r'\1<w:headerReference w:type="default" r:id="rId%d"/>'
             r'<w:footerReference w:type="default" r:id="rId%d"/>' % (hid, fid), doc, count=1)
CPAT = re.compile(r'<w:color w:val="(?:0F4761|156082|595959|272727)"[^>]*/>')
doc = CPAT.sub('<w:color w:val="000000"/>', doc).replace('>Table of Contents<', '>目\u3000录<')
parts['word/document.xml'] = doc.encode('utf-8')
parts['word/styles.xml'] = CPAT.sub('<w:color w:val="000000"/>',
                                    parts['word/styles.xml'].decode('utf-8')).encode('utf-8')

tmp = F + '.tmp'
zo = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
for k in names + [x for x in parts if x not in names]:
    zo.writestr(k, parts[k])
zo.close(); shutil.move(tmp, F)
print('已生成 %s（版本 %s）' % (F, REV))
